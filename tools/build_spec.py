"""Curated, verified field spec for FS26 ASES v1.0 -> templates/ases/fieldmap.json
Every field/table/image area is declared explicitly (sheet by sheet, in Excel order) and VERIFIED against the template:
 - field cells must be the top-left of their merge, must not hold a formula, and (unless noted) must be unlocked;
 - dropdown options are taken from the cell's own data validation;
 - labels/help come from the template's own text cells.
Run: python3 tools/build_spec.py   (prints a coverage/verification report)"""
import openpyxl, json, re, os, sys
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.utils.cell import range_boundaries
HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TPL = os.path.join(HERE, "templates/ases/FS26_ASES_v1.0.xlsm")
wb = openpyxl.load_workbook(TPL, keep_vba=True)
REVIEW = []  # items intentionally left out / uncertain

def sh(c, dc=0, dr=0):  # shift a cell ref
    m = re.match(r"([A-Z]+)(\d+)", c); return f"{L(CI(m.group(1))+dc)}{int(m.group(2))+dr}"
def sha(a, dc=0, dr=0): return ":".join(sh(x, dc, dr) for x in a.split(":"))
def txt(ws, c):
    v = ws[c].value
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None and not str(v).startswith("=") else ""

class Sheet:
    def __init__(s, name, title, intro_cells=(), guidance=()):
        s.ws = wb[name]; s.name = name; s.title = title; s.sections = []; s.fields = []; s.images = []; s.tables = []
        s.intro = " ".join(t for t in (txt(s.ws, c) for c in intro_cells) if t); s.guidance = list(guidance)
        s.dv = []
        for d in s.ws.data_validations.dataValidation:
            if d.type == "list" and d.formula1 and d.formula1.startswith('"'):
                s.dv.append((str(d.sqref).split(), [o.strip() for o in d.formula1.strip('"').split(",")]))
        s.merged = {}
        for m in s.ws.merged_cells.ranges:
            s.merged[f"{L(m.min_col)}{m.min_row}"] = str(m)
            for r in range(m.min_row, m.max_row+1):
                for c in range(m.min_col, m.max_col+1):
                    if (r, c) != (m.min_row, m.min_col): s.merged.setdefault("_inner", set()).add(f"{L(c)}{r}")
    def options(s, cell):
        col, row = re.match(r"([A-Z]+)(\d+)", cell).groups(); ci, ri = CI(col), int(row)
        for refs, opts in s.dv:
            for ref in refs:
                a, b, c2, d = range_boundaries(ref if ":" in ref else ref + ":" + ref)
                if a <= ci <= c2 and b <= ri <= d: return opts
        return None
    def section(s, title, help=""):
        s.cur = {"id": f"s{len(s.sections)+1}", "title": title, "help": help, "items": []}; s.sections.append(s.cur); return s.cur
    def field(s, cell, label, type=None, unit=None, help="", allow_locked=False, required=True):
        ws = s.ws; c = ws[cell]
        assert cell not in s.merged.get("_inner", set()), f"{s.name}!{cell} is not top-left of a merge"
        assert not (isinstance(c.value, str) and c.value.startswith("=")), f"{s.name}!{cell} holds a formula"
        if c.protection.locked is not False:
            assert allow_locked, f"{s.name}!{cell} is locked"
        opts = s.options(cell)
        t = type or ("select" if opts else "text")
        if t == "select": assert opts, f"{s.name}!{cell} select without validation list"
        f = {"cell": cell, "range": s.merged.get(cell, cell), "label": label, "type": t, "options": opts if t == "select" else None,
             "unit": unit, "default": None if c.value is None else str(c.value), "help": help, "section": s.cur["id"], "required": required}
        s.fields.append(f); s.cur["items"].append({"kind": "field", "cell": cell}); return f
    def image(s, sid, area, label, help=""):
        a, b, c2, d = range_boundaries(area); assert c2 >= a and d >= b
        s.images.append({"id": sid, "anchor": area.split(":")[0], "area": area, "label": label, "help": help, "section": s.cur["id"]})
        s.cur["items"].append({"kind": "image", "id": sid})
    def table(s, tid, start, cols, max_rows, label, help=""):
        s.tables.append({"id": tid, "start": start, "cols": cols, "max_rows": max_rows, "label": label, "help": help, "section": s.cur["id"]})
        s.cur["items"].append({"kind": "table", "id": tid})
        col, row = re.match(r"([A-Z]+)(\d+)", start).groups()
        for i in range(len(cols)):
            for r in (int(row), int(row)+max_rows-1):
                cc = f"{L(CI(col)+i)}{r}"; assert s.ws[cc].protection.locked is False, f"{s.name}!{cc} table cell locked"
    def out(s):
        return {"name": s.name, "title": s.title, "intro": s.intro, "guidance": s.guidance, "sections": s.sections,
                "fields": s.fields, "images": s.images, "tables": s.tables}

SHEETS = []
# ---------------- Cover Sheet ----------------
S = Sheet("Cover Sheet", "Cover Sheet & Compliance Matrix", ["A3", "A12"]); ws = S.ws
S.section("Team details")
for c, l in [("D15", "University Name"), ("O15", "Car No.(s) & Event(s)"), ("D16", "Team Contact"), ("N16", "Team Contact – Email Address"),
             ("D17", "Faculty Advisor"), ("N17", "Faculty Advisor – Email Address")]:
    S.field(c, l, "email" if "Email" in l else "text")
S.section("Compliance matrix", txt(ws, "B19") + " " + " ".join(t for t in (txt(ws, "A9")+": "+txt(ws, "C9"), txt(ws, "A10")+": "+txt(ws, "C10"))))
for r in [25, 27, 28, 30, 31, 32, 33, 34, 35, 36, 37, 38]:
    tab, rule = txt(ws, f"B{r}"), txt(ws, f"H{r}")
    proof = f"CAD proof: {txt(ws, f'P{r}')} · Datasheet proof: {txt(ws, f'R{r}')} · Manual calculation: {txt(ws, f'T{r}')}"
    if ws[f"N{r}"].protection.locked is False and f"N{r}" not in S.merged.get("_inner", set()):
        for col, ax in (("M", "X"), ("N", "Y"), ("O", "Z")):
            S.field(f"{col}{r}", f"{tab} ({rule}) – Compliance shown? Direction {ax}", help=proof)
    else:
        S.field(f"M{r}", f"{tab} ({rule}) – Compliance shown?", help=proof)
REVIEW.append("Cover Sheet!M39:O39 has a No/N/A/Yes dropdown but no row label in the template – left out.")
SHEETS.append(S)
# ---------------- Accumulator Pics ----------------
S = Sheet("Accumulator Pics", "Accumulator Pictures & Wall Build-up", ["A5"], ["EV5.5.3 Guidance Notes"]); ws = S.ws
S.section("Laminated structures (only if TSAC walls are composite)", txt(ws, "R2") + " Overall/skin/core thicknesses are taken automatically from the 3-point-bending tab.")
for r in range(8, 12):
    part = txt(ws, f"R{r}")
    for col, hdr in (("U", "U3"), ("V", "V3"), ("X", "X3"), ("Z", "Z3")):
        S.field(f"{col}{r}", f"{part} – {txt(ws, hdr)}", required=False)
S.section("Metallic structures", txt(ws, "R13"))
for r in range(14, 18):
    S.field(f"S{r}", f"{txt(ws, f'R{r}')} – Thickness (metallic)", "number", "mm", required=False)
for title, r0, dimlabel, imgs in [("Accumulator Floor", 12, "Width", ("B17", "B18")), ("External Wall", 36, "Height", ("B41", "B42")),
                                  ("Internal Wall", 60, "Height", ("B65", "B66")), ("Cover", 84, "Width", ("B89", "B90"))]:
    S.section(title, "Material 'Other' = composite: results come from the Alt. Materials tabs. Thickness is taken from the metallic/laminate tables above.")
    S.field(f"E{r0}", f"{title} – Material")
    S.field(f"E{r0+2}", f"{title} – {txt(ws, f'A{r0+2}').replace(' [mm]','')}", "number", "mm")
    top = r0 + 7; bot = {12: 33, 36: 57, 60: 81, 84: 105}[r0]
    S.image(imgs[0], f"A{top}:H{bot}", txt(ws, imgs[0]))
    S.image(imgs[1], f"I{top}:P{bot}", txt(ws, imgs[1]))
    if r0 == 12:
        S.section("Accumulator position in the chassis", txt(ws, "Q19"))
        S.field("V22", txt(ws, "Q22") + " – Image provided?"); S.field("V23", txt(ws, "Q23") + " – Image provided?")
        S.image("R25", "Q29:U41", txt(ws, "R25")); S.image("R27", "V29:Z41", txt(ws, "R27") + " " + txt(ws, "R28"))
    if r0 == 36:
        S.section("Additional pictures", txt(ws, "R45")); S.image("R45", "Q46:Z57", txt(ws, "Q42"))
REVIEW.append("Accumulator Pics!V61:W63 has a No/Yes dropdown in the template without a matching label – left out.")
SHEETS.append(S)
# ---------------- 3pt bending ----------------
S = Sheet("EV5.5.5 Alt. Matl - 3pt Bending", "Alt. Materials – 3-Point Bending Tests", [], ["EV5.5.3 Guidance Notes", "T3.5.9 Guidance Notes"]); ws = S.ws
for k, part in enumerate(["Floor", "Ext. Walls", "Int. Walls", "Cover"]):
    o = 19 * k
    S.section(f"{part} laminate panel", txt(ws, sh("A1", o)) + " – " + " ".join(txt(ws, sh(c, o)) for c in ("B9", "B11", "B27")))
    S.image(sh("B7", o), sha("B12:R18", o), f"{part}: " + txt(ws, sh("B7", o)))
    S.image(sh("B23", o), sha("B31:R42", o), f"{part}: " + txt(ws, sh("B23", o)), txt(ws, sh("B29", o)) + " " + txt(ws, sh("B30", o)))
    h = txt(ws, sh("A45", o))
    for c, lab, u in [("C46", "x1", "mm"), ("G46", "y1", "N"), ("C47", "x2", "mm"), ("G47", "y2", "N")]:
        S.field(sh(c, o), f"{part} – {lab} ({u})", "number", u, h)
    S.field(sh("C50", o), f"{part} – ymax (N)", "number", "N", txt(ws, sh("A49", o)))
    for c, lab in [("C53", "E53"), ("C54", "E54"), ("C55", "E55"), ("C56", "E56"), ("C57", "E57")]:
        sym = txt(ws, sh("B" + c[1:], o)) or txt(ws, sh("A" + c[1:], o))
        S.field(sh(c, o), f"{part} – {txt(ws, sh(lab, o))} {sym}", "number", "mm", txt(ws, sh("A52", o)))
    S.field(sh("C61", o), f"{part} – tx (mm)", "number", "mm", txt(ws, sh("E61", o)), required=False)
    S.table(sh("C68", o).replace(":", ""), sh("C68", o), [{"label": "F", "unit": "kN"}, {"label": "s", "unit": "mm"}], 300,
            f"{part}: logged test data (F, s)", " ".join(txt(ws, sh(c, o)) for c in ("A63", "A64", "A65")))
    S.image(sh("F70", o), sha("F72:R83", o), f"{part}: " + txt(ws, sh("F67", o)), txt(ws, sh("F70", o)) + " " + txt(ws, sh("F71", o)))
o = 76
S.section("Rig compliance test (optional, different layup)", txt(ws, "BY1") + " – " + txt(ws, "BY49"))
for c, lab, u in [("CA46", "x1", "mm"), ("CE46", "y1", "N"), ("CA47", "x2", "mm"), ("CE47", "y2", "N")]:
    S.field(c, f"Rig – {lab} ({u})", "number", u, txt(ws, "BY45"), required=False)
for c in ("CA51", "CA52", "CA53"):
    r = c[2:]; S.field(c, f"Rig – {txt(ws, 'CC'+r)} ({txt(ws, 'BZ'+r)})", "number", None if r == "51" else "mm", required=False)
S.table("CA68", "CA68", [{"label": "F", "unit": "kN"}, {"label": "s", "unit": "mm"}], 300, "Rig: logged test data (F, s)", txt(ws, "BY64"))
for c, area, lab in [("BZ7", "BZ12:CP18", "setup"), ("BZ23", "BZ31:CP42", "curve"), ("CD70", "CD72:CP83", "thickness")]:
    if txt(ws, c): S.image(c, area, "Rig: " + txt(ws, c))
    else: REVIEW.append(f"3pt Bending rig block: no template text at {c} ({lab} picture) – no upload slot offered.")
SHEETS.append(S)
# ---------------- Shear ----------------
S = Sheet("EV5.5.5 Alt. Matl - Shear", "Alt. Materials – Perimeter Shear Tests", [], ["EV5.5.3 Guidance Notes"]); ws = S.ws
for k, part in enumerate(["Floor", "Ext. Walls", "Int. Walls", "Cover"]):
    o = 19 * k
    S.section(f"{part} shear test", txt(ws, sh("A1", o)) + " – " + " ".join(txt(ws, sh(c, o)) for c in ("B9", "B11", "B27")))
    S.image(sh("B7", o), sha("B12:R18", o), f"{part}: " + txt(ws, sh("B7", o)))
    S.image(sh("B23", o), sha("B28:R44", o), f"{part}: " + txt(ws, sh("B23", o)), txt(ws, sh("B25", o)))
    S.field(sh("D48", o), f"{part} – First peak force (N) – determines shear strength", "number", "N", txt(ws, sh("A47", o)))
    S.field(sh("D49", o), f"{part} – Higher peak force (N)", "number", "N", txt(ws, sh("A47", o)))
    S.table(sh("C61", o), sh("C61", o), [{"label": "F", "unit": "kN"}, {"label": "s", "unit": "mm"}], 300,
            f"{part}: logged test data (F, s)", " ".join(txt(ws, sh(c, o)) for c in ("A56", "A57", "A58")))
    S.image(sh("F63", o), sha("F65:R76", o), f"{part}: " + txt(ws, sh("F60", o)), txt(ws, sh("F63", o)) + " " + txt(ws, sh("F64", o)))
SHEETS.append(S)
# ---------------- Datasheets ----------------
for name, title, spec in [
    ("Datasheets Fiber", "Datasheets – Fibre / Core Materials", [("A7", "Material type"), ("E9", "A9"), ("L9", "H9")]),
    ("Datasheets Other Matl", "Datasheets – Other Materials", [("A7", "Material type"), ("E9", "A9"), ("L9", "H9"), ("E11", "A11"), ("L11", "H11")]),
    ("Datasheets Glue", "Datasheets – Adhesives", [("E7", "A7"), ("L7", "H7"), ("E9", "A9"), ("L9", "H9")])]:
    S = Sheet(name, title, [], ["EV5.5.3 Guidance Notes"]); ws = S.ws
    first_img_row = {"Datasheets Fiber": 17, "Datasheets Other Matl": 19, "Datasheets Glue": 17}[name]
    for k in range(8):
        o = 17 * k
        S.section(f"Material {k+1}", "Fill one block per material. Leave unused blocks empty.")
        S.field(sh("A5", o), f"Material {k+1} – Name / designation", required=(k == 0))
        for c, lab in spec:
            lt = lab if " " in lab else txt(ws, sh(lab, o))
            is_num = "[MPa]" in lt
            S.field(sh(c, o), f"Material {k+1} – {lt.replace('  ', ' ')}", "number" if is_num else None, "MPa" if is_num else None, required=(k == 0))
        hl = " ".join(txt(ws, sh(f"A{first_img_row+i}", o)) for i in range(3))
        S.image(sh(f"A{first_img_row+3}", o), sha(f"A{first_img_row+3}:P118", o), f"Material {k+1}: datasheet pages", hl)
    SHEETS.append(S)
# ---------------- Cell fixation ----------------
S = Sheet("EV5.5.10-12 Cell Fixation", "Cell Fixation", [], ["EV5.5.10-11 Guidance Notes", "EV5.5.12 Guidance Notes"]); ws = S.ws
S.section("Cell & fixation data", "Shear/compression stresses and PASS/FAIL are calculated by the spreadsheet (2027: EV5.5.14 / EV5.5.15).")
for c, t, u in [("M12", None, None), ("M13", None, None), ("M14", "text", None), ("M15", "number", "g"), ("M16", "number", "mm"), ("M17", "number", "mm"),
                ("M20", "number", None), ("M22", "number", "N"), ("M25", "number", None)]:
    lab = txt(ws, "A" + c[1:]); lab = re.sub(r"\s*/\s*(g|mm)$", "", lab)
    S.field(c, lab, t, u, required=c not in ("M14", "M22", "M25"))
S.field("M28", txt(ws, "A28"), allow_locked=True, help="Template note: if 'Yes', explain how EV4.5.12 is fulfilled; cannot be N/A for pouch cells.")
REVIEW.append("Cell Fixation!M28 (compressible materials) is LOCKED in the template although it has a dropdown – included (the app writes it), please confirm FSG intends it as input.")
S.section("Pictures & calculations", txt(ws, "B6"))
S.image("B6", "S5:AK30", txt(ws, "B6"))
S.image("B31", "S31:AK60", "Fixation concept, calculations, friction test", " ".join(txt(ws, f"B{r}") for r in range(31, 46)))
REVIEW.append("Cell Fixation: the whole tab is an unlocked free canvas; pictures are placed to the right of the form (S5:AK30 and S31:AK60) – check placement.")
SHEETS.append(S)
# ---------------- EV5.5.9 Segment / TSAC / attachments ----------------
for name, title in [("EV5.5.9 Segment", "Segment (cell mounting within segment)"), ("EV5.5.9 TSAC", "TSAC"), ("EV5.5.9 TSAC attachments", "TSAC Attachments")]:
    S = Sheet(name, title, ["B6", "B7", "B8", "B9", "B10"], ["EV5.5.9 Guidance Notes"] + (["EV5.5 TSAC Attachment Guidance", "T3.5.9 Guidance Notes"] if "attach" in name else [])); ws = S.ws
    S.section("Materials", "")
    S.field("AF2", txt(ws, "S2"), help="If Yes: show mechanical properties at 60 °C in the material datasheets.")
    S.field("AF6", txt(ws, "S6"), help="If Yes: show mechanical properties at 60 °C in the material datasheets.")
    S.section("Load cases (ISO 8855 vehicle coordinates)", "CAD/FEA pictures, calculations and explanations per direction.")
    for col, area in (("A", "A13:P116"), ("R", "R13:AG116"), ("AI", "AI13:AX116")):
        S.image(f"{col}12", area, txt(ws, f"{col}12"))
    SHEETS.append(S)
# ---------------- Fasteners ----------------
S = Sheet("EV5.5.13 Fasteners", "Fasteners", ["B6", "B7", "B8"], []); ws = S.ws
n = 0
for r0 in (11, 35, 59, 83):
    for col, fc in (("A", "H"), ("R", "Y"), ("AI", "AP")):
        n += 1; o = CI(fc) - CI("H"); req = n == 1
        S.section(f"Fastener group {n}", "Group fasteners of the same size/application. Leave unused groups empty.")
        S.field(f"{fc}{r0}", f"Group {n} – {txt(ws, f'{col}{r0}')}", required=req)
        S.field(f"{fc}{r0+1}", f"Group {n} – {txt(ws, f'{col}{r0+1}')}", required=req)
        S.field(f"{fc}{r0+2}", f"Group {n} – {txt(ws, f'{col}{r0+2}')}", "text", required=req)
        S.image(f"{col}{r0+3}", sha(f"A{r0+4}:P{r0+22}", CI(col)-1), f"Group {n}: " + txt(ws, f"{col}{r0+3}"))
REVIEW.append("Fasteners: dropdowns at F10:G10 and AF2/AF6 exist in the template without labels – left out.")
SHEETS.append(S)
# ---------------- Openings ----------------
S = Sheet("EV5.5.14-17 Openings&Separation", "Openings & Segment Separation", [], []); ws = S.ws
S.section("Segments", "2027: limits 120 V / 6 MJ / 12 kg apply per mechanical section (EV5.3.4).")
for r in range(5, 21):
    seg = txt(ws, f"B{r}")
    S.field(f"E{r}", f"{seg} – Mass", "number", "kg", required=r == 5); S.field(f"G{r}", f"{seg} – Voltage", "number", "V", required=r == 5)
    S.field(f"I{r}", f"{seg} – Energy", "number", "MJ", required=r == 5); S.field(f"K{r}", f"{seg} – Separation shown?", required=r == 5)
S.section("Openings, ducts & separation", txt(ws, "A22"))
for r in (24, 26, 27, 28, 29, 30, 31, 33): S.field(f"R{r}", txt(ws, f"A{r}"))
S.image("A35", "A35:S62", "Pictures / CAD with dimensions & labels")
REVIEW.append("Openings: the tab is an unlocked free canvas without a marked picture area; pictures are placed below the checklist (A35:S62) – check placement.")
SHEETS.append(S)
# ---------------- Additional Info ----------------
S = Sheet("Additional Info", "Additional Information", ["A5"], []); ws = S.ws
S.section("Additional information", "Free text (one paragraph per line) is written from B9 downwards; N/A justifications from all tabs are appended automatically.")
S.field("B9", "Additional information text", "textarea", required=False, allow_locked=True)
S.image("B30", "B30:S63", "Additional pictures / CAD / calculations")
SHEETS.append(S)
# ---------------- Version History / Rules clarifications ----------------
S = Sheet("Version History", "Version History", ["A1"], []); ws = S.ws
S.section("Submission versions"); S.table("A5", "A5", [{"label": txt(ws, "A4")}, {"label": txt(ws, "B4")}], 114, "Version history")
SHEETS.append(S)
S = Sheet("Rules Clarifications", "Rules Clarifications", ["A1"], []); ws = S.ws
S.section("Clarifications"); S.table("A4", "A4", [{"label": txt(ws, "A3")}, {"label": txt(ws, "B3")}], 174, "Rules clarifications")
SHEETS.append(S)

# guidance texts
G = {}
for gname in ["Read_Me", "T3.5.9 Guidance Notes", "EV5.5.3 Guidance Notes", "EV5.5.9 Guidance Notes", "EV5.5.10-11 Guidance Notes", "EV5.5.12 Guidance Notes", "EV5.5 TSAC Attachment Guidance"]:
    w = wb[gname]; G[gname] = "\n".join(str(c.value).strip() for row in w.iter_rows() for c in row if isinstance(c.value, str) and c.value.strip() and not c.value.startswith("="))
out = {"template": os.path.basename(TPL), "sheets": [s.out() for s in SHEETS], "guidance": G, "needs_review": REVIEW}
json.dump(out, open(os.path.join(HERE, "templates/ases/fieldmap.json"), "w"), ensure_ascii=False, indent=1)
tot = [0, 0, 0]
for s in SHEETS:
    print(f"{s.name:36} fields={len(s.fields):3} images={len(s.images):3} tables={len(s.tables)}"); tot[0] += len(s.fields); tot[1] += len(s.images); tot[2] += len(s.tables)
print("TOTAL", tot); print("NEEDS REVIEW:"); [print(" -", r) for r in REVIEW]
