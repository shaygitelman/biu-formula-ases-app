import openpyxl, json, sys
from openpyxl.utils import get_column_letter as L
wb = openpyxl.load_workbook('templates/ases/FS26_ASES_v1.0.xlsm', keep_vba=True)
fm = json.load(open('templates/ases/fieldmap.json')); mapped = {s['name']: {f['cell'] for f in s['fields']} for s in fm['sheets']}
out = {}
for ws in wb:
    if ws.sheet_state != 'visible' or 'Guidance' in ws.title or ws.title=='Read_Me': continue
    tl = {}
    for mr in ws.merged_cells.ranges:
        for r in range(mr.min_row, mr.max_row+1):
            for c in range(mr.min_col, mr.max_col+1): tl[(r,c)] = (mr.min_row, mr.min_col)
    U = set()
    for row in ws.iter_rows():
        for c in row:
            if c.protection.locked is False: U.add((c.row, c.column))
    seen=set(); comps=[]
    for p in sorted(U):
        if p in seen: continue
        st=[p]; comp=[]; seen.add(p)
        while st:
            q=st.pop(); comp.append(q)
            for d in ((0,1),(1,0),(0,-1),(-1,0)):
                n=(q[0]+d[0],q[1]+d[1])
                if n in U and n not in seen: seen.add(n); st.append(n)
        nodes={tl.get(x,x) for x in comp}
        comps.append((nodes, comp))
    print(f"== {ws.title} comps={len(comps)}")
    for nodes, comp in comps:
        rs=[x[0] for x in comp]; cs=[x[1] for x in comp]
        bb=f"{L(min(cs))}{min(rs)}:{L(max(cs))}{max(rs)}"
        if len(nodes)>12:
            print(f"  BIG {bb} nodes={len(nodes)}")
        else:
            for (r,c) in sorted(nodes):
                co=f"{L(c)}{r}"; v=ws.cell(r,c).value
                print(f"  {co:6} {'M' if co in mapped.get(ws.title,()) else '-'} val={str(v)[:30]!r}")
