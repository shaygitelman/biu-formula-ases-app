import openpyxl, json, sys
from openpyxl.utils import get_column_letter as L
wb = openpyxl.load_workbook('templates/ases/FS26_ASES_v1.0.xlsm', keep_vba=True)
fm = json.load(open('templates/ases/fieldmap.json')); mapped = {s['name']: {f['cell'] for f in s['fields']} for s in fm['sheets']}
for ws in wb:
    if ws.sheet_state != 'visible': continue
    tl = {}
    for mr in ws.merged_cells.ranges:
        for r in range(mr.min_row, mr.max_row+1):
            for c in range(mr.min_col, mr.max_col+1): tl[(r,c)] = (mr.min_row, mr.min_col)
    un = []
    for row in ws.iter_rows():
        for c in row:
            if c.protection.locked is False and tl.get((c.row,c.column),(c.row,c.column)) == (c.row,c.column):
                un.append(c.coordinate)
    dv = [(str(d.sqref), d.type, d.formula1) for d in ws.data_validations.dataValidation]
    m = mapped.get(ws.title, set())
    miss = [u for u in un if u not in m]
    print(f"== {ws.title}: unlocked={len(un)} mapped={len(m)} unmapped={len(miss)} imgs={len(ws._images)} dv={len(dv)}")
    if miss and ws.title in mapped: print("   UNMAPPED:", " ".join(miss[:80]))
    if ws.title not in mapped and un: print("   (sheet not mapped) unlocked:", " ".join(un[:30]))
