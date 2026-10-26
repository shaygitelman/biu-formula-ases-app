"""BIU Formula – FSG document filler (multi-template). Templates live in templates/<id>/ (template file, fieldmap.json, rule_map.json)."""
import os, json, sqlite3, time, re, io, base64, hashlib
from flask import Flask, request, jsonify, send_file, send_from_directory, g, abort, Response
import exporter, checks
HERE = os.path.dirname(os.path.abspath(__file__)); TD = os.path.join(HERE, "templates")
REG = json.load(open(os.path.join(TD, "registry.json")))
RULES = json.load(open(os.path.join(HERE, "data/rules_2027.json")))
T = {}
for tid, meta in REG.items():
    fmp = os.path.join(TD, tid, "fieldmap.json")
    if not meta.get("enabled") or not os.path.exists(fmp): continue
    fm = json.load(open(fmp)); rmp = os.path.join(TD, tid, "rule_map.json")
    S = {s["name"]: s for s in fm["sheets"]}
    for s in fm["sheets"]:
        s["_f"] = {f["cell"]: f for f in s["fields"]}; s["_i"] = {i["id"]: i for i in s["images"]}; s["_t"] = {t["id"]: t for t in s.get("tables", [])}
    T[tid] = {"meta": meta, "fm": fm, "sheets": S, "file": os.path.join(TD, tid, meta["file"]),
              "rule_map": json.load(open(rmp)) if os.path.exists(rmp) else {}}
DD = os.environ.get("DATA_DIR", os.path.join(HERE, "data")); os.makedirs(DD, exist_ok=True)
DB = os.environ.get("DB_PATH", os.path.join(DD, "app.db"))
UP = os.path.join(DD, "uploads") if os.environ.get("DATA_DIR") else os.path.join(HERE, "uploads"); os.makedirs(UP, exist_ok=True)
app = Flask(__name__, static_folder="static"); app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024
ACCESS = os.environ.get("ACCESS_CODE")
IMG_RE = re.compile(r"\.(png|jpe?g|gif|bmp|webp)$", re.I)

@app.before_request
def gate():
    if not ACCESS: return
    a = request.authorization
    if not a or a.password != ACCESS: return Response("Access code required", 401, {"WWW-Authenticate": 'Basic realm="BIU ASES"'})
def tpl(tid):
    if tid not in T: abort(404)
    return T[tid]
def sheet_of(tid, name):
    s = tpl(tid)["sheets"].get(name)
    if s is None: abort(404)
    return s
def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB); g.db.row_factory = sqlite3.Row
        g.db.executescript("""create table if not exists vals(tpl text, sheet text, cell text, value text, user text, ts real, primary key(tpl,sheet,cell));
          create table if not exists imgs(id integer primary key, tpl text, sheet text, slot text, anchor text, file text, name text, user text, ts real);
          create table if not exists notes(tpl text, sheet text, text text, user text, ts real, primary key(tpl,sheet));
          create table if not exists flags(tpl text, sheet text, section text, skip integer, primary key(tpl,sheet,section));
          create table if not exists log(ts real, tpl text, user text, sheet text, cell text, value text);""")
    return g.db
@app.teardown_appcontext
def close(e):
    d = g.pop("db", None)
    if d: d.commit(); d.close()

def state(tid, name):
    d = db()
    vals = {r["cell"]: r["value"] for r in d.execute("select cell,value from vals where tpl=? and sheet=?", (tid, name))}
    imgs = [dict(r) for r in d.execute("select id,slot,name,user,ts from imgs where tpl=? and sheet=? order by ts", (tid, name))]
    n = d.execute("select text from notes where tpl=? and sheet=?", (tid, name)).fetchone()
    skip = {r["section"] for r in d.execute("select section from flags where tpl=? and sheet=? and skip=1", (tid, name))}
    return vals, imgs, (n["text"] if n else ""), skip
def counts(imgs):
    c = {}
    for i in imgs: c[i["slot"]] = c.get(i["slot"], 0) + 1
    return c
def progress(tid, name, st=None):
    s = sheet_of(tid, name); vals, imgs, notes, skip = st or state(tid, name); cnt = counts(imgs)
    if "*" in skip: return {"done": 0, "total": 0, "pct": 100, "na": True, "errors": 0}
    act = checks.active_sections(s, vals, cnt, skip); done = tot = 0; err = 0
    for f in s["fields"]:
        v = vals.get(f["cell"])
        if checks.field_hint(f, v): err += 1
        if f["section"] in act and (f.get("required") or v not in (None, "")): tot += 1; done += v not in (None, "")
    for i in s["images"]:
        if i["section"] in act: tot += 1; done += bool(cnt.get(i["id"]))
    for t in s.get("tables", []):
        if t["section"] in act: tot += 1; done += bool(vals.get("T:" + t["id"]))
    return {"done": done, "total": tot, "pct": round(100*done/tot) if tot else 100, "na": False, "optional": not tot, "errors": err}
def findings(tid, name):
    t = tpl(tid); vals, imgs, notes, skip = state(tid, name)
    if "*" in skip: return [{"level": "info", "msg": "Tab marked as not applicable – it will be exported unchanged."}], vals
    return checks.review(t["sheets"][name], vals, counts(imgs), notes, skip) + checks.rule_findings(name, t["rule_map"], RULES), vals
def rules_ctx(tid, name):
    rm = tpl(tid)["rule_map"].get(name) or {}
    return "\n".join(f"{r}: {RULES[r]}" for r in rm.get("rules_2027", []) if r in RULES) + (f"\nNUMBERING NOTE: {rm['mismatch']}" if rm.get("mismatch") else "")

@app.get("/")
def index(): return send_from_directory(HERE + "/static", "index.html")
@app.get("/api/templates")
def templates(): return jsonify({k: {**v, "loaded": k in T} for k, v in REG.items()})
@app.get("/api/rules/<rid>")
def rule(rid): r = RULES.get(rid.replace(" ", "")); return jsonify({"id": rid, "text": r}) if r else (jsonify({"error": "not found"}), 404)
@app.get("/api/<tid>/fieldmap")
def fieldmap(tid):
    t = tpl(tid); fm = t["fm"]
    return jsonify({"template": fm["template"], "title": t["meta"]["title"], "guidance": fm.get("guidance", {}), "needs_review": fm.get("needs_review", []),
        "sheets": [{**{k: v for k, v in s.items() if not k.startswith("_")}, "rules": t["rule_map"].get(s["name"], {}),
                    "rule_texts": {r: RULES[r] for r in t["rule_map"].get(s["name"], {}).get("rules_2027", []) if r in RULES}} for s in fm["sheets"]]})
@app.get("/api/<tid>/progress")
def all_progress(tid): return jsonify({n: progress(tid, n) for n in tpl(tid)["sheets"]})
@app.get("/api/<tid>/overview")
def overview(tid):
    out = []
    for n in tpl(tid)["sheets"]:
        p = progress(tid, n); f, _ = findings(tid, n)
        out.append({"sheet": n, **p, "warn": sum(x["level"] == "warn" for x in f), "error": sum(x["level"] == "error" for x in f)})
    return jsonify(out)
@app.get("/api/<tid>/sheet/<path:name>")
def get_sheet(tid, name):
    st = state(tid, sheet_of(tid, name)["name"]); vals, imgs, notes, skip = st
    return jsonify({"values": vals, "images": imgs, "notes": notes, "skip": sorted(skip), "progress": progress(tid, name, st)})
def put_value(tid, sheet, cell, value, user):
    s = sheet_of(tid, sheet)
    if cell.startswith("T:"):
        t = s["_t"].get(cell[2:])
        if t is None: return "unknown table"
        if value not in (None, ""):
            rows = json.loads(value) if isinstance(value, str) else value
            if not isinstance(rows, list) or len(rows) > t["max_rows"]: return f"table accepts at most {t['max_rows']} rows"
            rows = [[("" if c is None else str(c)) for c in (r + [""] * len(t["cols"]))[:len(t["cols"])]] for r in rows if isinstance(r, list)]
            while rows and not any(rows[-1]): rows.pop()
            value = json.dumps(rows) if rows else None
    elif cell not in s["_f"]: return "unknown field"
    d = db()
    if value in (None, ""): d.execute("delete from vals where tpl=? and sheet=? and cell=?", (tid, sheet, cell))
    else: d.execute("insert or replace into vals values(?,?,?,?,?,?)", (tid, sheet, cell, str(value), user, time.time()))
    d.execute("insert into log values(?,?,?,?,?,?)", (time.time(), tid, user, sheet, cell, None if value is None else str(value)[:2000]))
@app.post("/api/<tid>/value")
def set_value(tid):
    j = request.json; err = put_value(tid, j["sheet"], j["cell"], j.get("value"), j.get("user") or "anon")
    if err: return jsonify({"error": err}), 400
    f = sheet_of(tid, j["sheet"])["_f"].get(j["cell"])
    return jsonify({"ok": True, "hint": checks.field_hint(f, j.get("value")) if f else None, "progress": progress(tid, j["sheet"])})
@app.post("/api/<tid>/skip")
def set_skip(tid):
    j = request.json; s = sheet_of(tid, j["sheet"])
    if j["section"] != "*" and j["section"] not in {x["id"] for x in s["sections"]}: return jsonify({"error": "unknown section"}), 400
    db().execute("insert or replace into flags values(?,?,?,?)", (tid, j["sheet"], j["section"], 1 if j.get("skip") else 0))
    return jsonify({"ok": True, "progress": progress(tid, j["sheet"])})
@app.post("/api/<tid>/notes")
def set_notes(tid):
    j = request.json; sheet_of(tid, j["sheet"])
    db().execute("insert or replace into notes values(?,?,?,?,?)", (tid, j["sheet"], j.get("text", ""), j.get("user"), time.time())); return jsonify({"ok": True})
@app.post("/api/<tid>/image")
def up_image(tid):
    sheet, slot = request.form["sheet"], request.form["slot"]; s = sheet_of(tid, sheet)
    if slot not in s["_i"]: return jsonify({"error": "unknown picture slot"}), 400
    added = []
    for f in request.files.getlist("file"):
        if not IMG_RE.search(f.filename or ""): return jsonify({"error": f"{f.filename}: only image files (png, jpg, gif, bmp, webp)"}), 400
        data = f.read()
        try:
            from PIL import Image; Image.open(io.BytesIO(data)).verify()
        except Exception: return jsonify({"error": f"{f.filename}: not a readable image"}), 400
        fn = f"{int(time.time()*1000)}_{hashlib.md5(data).hexdigest()[:8]}_{re.sub(r'[^A-Za-z0-9._-]', '_', f.filename)[-60:]}"
        open(os.path.join(UP, fn), "wb").write(data)
        cur = db().execute("insert into imgs(tpl,sheet,slot,anchor,file,name,user,ts) values(?,?,?,?,?,?,?,?)", (tid, sheet, slot, s["_i"][slot]["anchor"], fn, f.filename, request.form.get("user"), time.time()))
        added.append(cur.lastrowid)
    return jsonify({"ok": True, "ids": added, "progress": progress(tid, sheet)})
@app.delete("/api/<tid>/image/<int:iid>")
def del_image(tid, iid): db().execute("delete from imgs where id=? and tpl=?", (iid, tid)); return jsonify({"ok": True})
@app.get("/api/<tid>/image/<int:iid>")
def get_image(tid, iid):
    r = db().execute("select file from imgs where id=? and tpl=?", (iid, tid)).fetchone()
    if not r: abort(404)
    if request.args.get("thumb"):
        from PIL import Image
        p = os.path.join(UP, r["file"]); tp = p + ".thumb.jpg"
        if not os.path.exists(tp):
            im = Image.open(p); im.thumbnail((360, 360)); im.convert("RGB").save(tp, "JPEG", quality=80)
        return send_file(tp, max_age=86400)
    return send_from_directory(UP, r["file"])
@app.get("/api/<tid>/check/<path:name>")
def check(tid, name): sheet_of(tid, name); return jsonify(findings(tid, name)[0])
@app.post("/api/<tid>/chat")
def chat(tid):
    j = request.json; name = j["sheet"]; s = sheet_of(tid, name); f, vals = findings(tid, name)
    guid = "\n".join(tpl(tid)["fm"]["guidance"].get(k, "") for k in s.get("guidance", []) + ["Read_Me"])
    return jsonify({"answer": checks.assistant(s, j["message"], vals, [x for x in f if x["level"] != "rule"], guid, j.get("history"), rules_ctx(tid, name)), "llm": bool(checks.llm_config())})

# ---------- export ----------
def col_shift(cell, k):
    c, r = exporter.split(cell); return exporter.n2col(exporter.col2n(c) + k) + str(r)
def row_shift(cell, k):
    c, r = exporter.split(cell); return c + str(r + k)
def build_export(tid, path):
    d = db(); t = tpl(tid); S = t["sheets"]; values, images = {}, {}
    skipped = {(r["sheet"], r["section"]) for r in d.execute("select sheet,section from flags where tpl=? and skip=1", (tid,))}
    for r in d.execute("select sheet,cell,value from vals where tpl=?", (tid,)):
        s = S.get(r["sheet"])
        if not s or (r["sheet"], "*") in skipped: continue
        if r["cell"].startswith("T:"):
            tb = s["_t"].get(r["cell"][2:])
            if not tb: continue
            for i, row in enumerate(json.loads(r["value"])[:tb["max_rows"]]):
                for k, v in enumerate(row):
                    if v not in (None, ""): values.setdefault(r["sheet"], {})[col_shift(row_shift(tb["start"], i), k)] = v
        elif r["cell"] in s["_f"]:
            f = s["_f"][r["cell"]]; v = r["value"]
            if f["type"] == "number" and checks.num(v) is not None: v = checks.num(v)
            if r["cell"] == "B9" and r["sheet"] == "Additional Info": continue
            values.setdefault(r["sheet"], {})[r["cell"]] = v
    for r in d.execute("select sheet,slot,file,name from imgs where tpl=? order by ts", (tid,)):
        s = S.get(r["sheet"]); sl = s and s["_i"].get(r["slot"])
        if not sl or (r["sheet"], "*") in skipped: continue
        images.setdefault(r["sheet"], []).append((sl.get("area") or sl["anchor"], open(os.path.join(UP, r["file"]), "rb").read(), r["name"]))
    if "Additional Info" in S:  # free text + N/A justifications -> B9.. (one line per row, above the picture area at B30)
        txt = (d.execute("select value from vals where tpl=? and sheet='Additional Info' and cell='B9'", (tid,)).fetchone() or {"value": ""})["value"] or ""
        lines = [l for l in txt.split("\n")]
        for n in d.execute("select sheet,text from notes where tpl=? and text<>'' order by sheet", (tid,)):
            if (n["sheet"], "*") not in skipped: lines += [f"[{n['sheet']}] {l}" for l in n["text"].split("\n") if l.strip()]
        while lines and not lines[-1].strip(): lines.pop()
        if len(lines) > 21: lines = lines[:20] + [" / ".join(lines[20:])]
        for i, l in enumerate(lines):
            if l.strip(): values.setdefault("Additional Info", {})[f"B{9+i}"] = l
    return exporter.export(t["file"], path, values, images)
@app.get("/api/<tid>/export")
def export(tid):
    ext = os.path.splitext(tpl(tid)["file"])[1]; path = os.path.join(DD, f"export_{tid}{ext}"); build_export(tid, path)
    return send_file(path, as_attachment=True, download_name=f"BIU_{tid.upper()}_{time.strftime('%Y%m%d_%H%M')}{ext}")

# ---------- backup / restore / import ----------
@app.get("/api/<tid>/backup")
def backup(tid):
    d = db(); tpl(tid)
    out = {"format": "biu-ases-backup/1", "template": tid, "created": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
           "values": [dict(r) for r in d.execute("select sheet,cell,value,user,ts from vals where tpl=?", (tid,))],
           "notes": [dict(r) for r in d.execute("select sheet,text,user,ts from notes where tpl=?", (tid,))],
           "flags": [dict(r) for r in d.execute("select sheet,section,skip from flags where tpl=?", (tid,))], "images": []}
    for r in d.execute("select sheet,slot,file,name,user,ts from imgs where tpl=? order by ts", (tid,)):
        p = os.path.join(UP, r["file"])
        if os.path.exists(p): out["images"].append({**{k: r[k] for k in ("sheet", "slot", "name", "user", "ts")}, "data": base64.b64encode(open(p, "rb").read()).decode()})
    return Response(json.dumps(out), mimetype="application/json", headers={"Content-Disposition": f'attachment; filename="BIU_{tid.upper()}_backup_{time.strftime("%Y%m%d_%H%M")}.json"'})
@app.post("/api/<tid>/restore")
def restore(tid):
    S = tpl(tid)["sheets"]
    try: j = json.load(request.files["file"])
    except Exception: return jsonify({"error": "not a valid backup file"}), 400
    if j.get("format") != "biu-ases-backup/1" or j.get("template") != tid: return jsonify({"error": "this backup is not for this template"}), 400
    d = db(); bk = os.path.join(DD, f"auto_before_restore_{int(time.time())}.json")
    open(bk, "w").write(backup(tid).get_data(as_text=True))  # safety copy of the current state
    for tb in ("vals", "imgs", "notes", "flags"): d.execute(f"delete from {tb} where tpl=?", (tid,))
    n = {"values": 0, "images": 0, "skipped": 0}
    for v in j.get("values", []):
        if v["sheet"] in S and (v["cell"] in S[v["sheet"]]["_f"] or v["cell"][2:] in S[v["sheet"]]["_t"]):
            d.execute("insert or replace into vals values(?,?,?,?,?,?)", (tid, v["sheet"], v["cell"], v["value"], v.get("user"), v.get("ts") or time.time())); n["values"] += 1
        else: n["skipped"] += 1
    for v in j.get("notes", []):
        if v["sheet"] in S: d.execute("insert or replace into notes values(?,?,?,?,?)", (tid, v["sheet"], v["text"], v.get("user"), v.get("ts")))
    for v in j.get("flags", []):
        if v["sheet"] in S: d.execute("insert or replace into flags values(?,?,?,?)", (tid, v["sheet"], v["section"], v["skip"]))
    for v in j.get("images", []):
        s = S.get(v["sheet"])
        if not s or v["slot"] not in s["_i"]: n["skipped"] += 1; continue
        data = base64.b64decode(v["data"]); fn = f"{int(time.time()*1000)}_{hashlib.md5(data).hexdigest()[:8]}_{re.sub(r'[^A-Za-z0-9._-]', '_', v['name'])[-60:]}"
        open(os.path.join(UP, fn), "wb").write(data)
        d.execute("insert into imgs(tpl,sheet,slot,anchor,file,name,user,ts) values(?,?,?,?,?,?,?,?)", (tid, v["sheet"], v["slot"], s["_i"][v["slot"]]["anchor"], fn, v["name"], v.get("user"), v.get("ts") or time.time())); n["images"] += 1
    return jsonify({"ok": True, **n})
@app.post("/api/<tid>/import")
def import_xlsm(tid):
    """Import typed values from an already filled copy of the same template. Pictures are NOT imported."""
    import openpyxl
    S = tpl(tid)["sheets"]; f = request.files["file"]; mode = request.form.get("mode", "fill")  # fill = only empty fields, overwrite = all
    try: wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=False, keep_vba=False)
    except Exception: return jsonify({"error": "could not read this file as Excel"}), 400
    ref = openpyxl.load_workbook(tpl(tid)["file"], data_only=False)
    miss = [n for n in S if n not in wb.sheetnames]
    if len(miss) > 3: return jsonify({"error": "this file does not look like the same template (missing tabs: " + ", ".join(miss[:5]) + ")"}), 400
    d = db(); n = 0; kept = 0
    def norm(v):
        if v is None: return None
        if isinstance(v, float) and v.is_integer(): v = int(v)
        v = str(v).strip(); return v or None
    for name, s in S.items():
        if name not in wb.sheetnames: continue
        ws, rs = wb[name], ref[name]; cur, _, _, _ = state(tid, name)
        items = [(f["cell"], norm(ws[f["cell"]].value), norm(rs[f["cell"]].value)) for f in s["fields"]]
        for c, v, dflt in items:
            if v is None or v == dflt or str(v).startswith("="): continue
            if c == "B9" and name == "Additional Info":
                lines = [norm(ws[f"B{r}"].value) for r in range(9, 30)]; v = "\n".join(l or "" for l in lines).strip()
            if cur.get(c) not in (None, "") and mode != "overwrite": kept += 1; continue
            put_value(tid, name, c, v, "import"); n += 1
        for tb in s.get("tables", []):
            rows = []
            for i in range(tb["max_rows"]):
                rows.append([norm(ws[col_shift(row_shift(tb["start"], i), k)].value) or "" for k in range(len(tb["cols"]))])
            ref_rows = [[norm(rs[col_shift(row_shift(tb["start"], i), k)].value) or "" for k in range(len(tb["cols"]))] for i in range(tb["max_rows"])]
            if rows == ref_rows: continue
            if cur.get("T:" + tb["id"]) and mode != "overwrite": kept += 1; continue
            put_value(tid, name, "T:" + tb["id"], json.dumps(rows), "import"); n += 1
    return jsonify({"ok": True, "imported": n, "kept_existing": kept, "note": "Typed values were imported. Pictures cannot be imported reliably – please upload them again."})
if __name__ == "__main__":
    app.run(host=os.environ.get("HOST", "127.0.0.1"), port=int(os.environ.get("PORT", 8077)), debug=False)
